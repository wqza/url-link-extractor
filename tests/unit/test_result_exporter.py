"""ResultExporter 单元测试（T13.02）。"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from url_link_extractor.core.result_exporter import ExcelResultExporter
from url_link_extractor.exceptions import ExcelWriteError
from url_link_extractor.models.enums import ErrorCode, RecordStatus
from url_link_extractor.models.export_spec import EXPORT_FILE_SPEC
from url_link_extractor.models.result import LinkTitleRecord


@pytest.fixture
def exporter():
    return ExcelResultExporter()


def _record(url, title="T", status=RecordStatus.SUCCESS):
    return LinkTitleRecord(url=url, title=title, status=status)


class TestExport:
    async def test_normal_export(self, exporter, tmp_path):
        records = [_record("https://e.com/a.html", "标题A"), _record("https://e.com/b.html", "标题B")]
        outcome = exporter.export(records, work_dir=tmp_path)
        assert outcome.error is None
        assert outcome.file_path == str(tmp_path / "result.xlsx")
        wb = load_workbook(outcome.file_path)
        ws = wb.active
        assert ws.cell(1, 1).value == "URL 完整链接"
        assert ws.cell(2, 1).value == "https://e.com/a.html"

    async def test_empty_result_only_header(self, exporter, tmp_path):
        outcome = exporter.export([], work_dir=tmp_path)
        assert outcome.error is None
        wb = load_workbook(outcome.file_path)
        ws = wb.active
        assert ws.max_row == 1

    async def test_title_none_empty_string(self, exporter, tmp_path):
        records = [_record("https://e.com/a.html", title=None, status=RecordStatus.NO_TITLE)]
        outcome = exporter.export(records, work_dir=tmp_path)
        wb = load_workbook(outcome.file_path)
        ws = wb.active
        assert ws.cell(2, 2).value in (None, "")

    async def test_overwrite_old_file(self, exporter, tmp_path):
        exporter.export([_record("https://e.com/old.html", "旧")], work_dir=tmp_path)
        exporter.export([_record("https://e.com/new.html", "新")], work_dir=tmp_path)
        wb = load_workbook(str(tmp_path / "result.xlsx"))
        ws = wb.active
        assert ws.cell(2, 1).value == "https://e.com/new.html"
        assert ws.max_row == 2

    async def test_export_failure_soft(self, tmp_path):
        # mock writer 抛异常
        from unittest.mock import MagicMock

        writer = MagicMock()
        writer.write = MagicMock(side_effect=ExcelWriteError("disk full"))
        exporter = ExcelResultExporter(writer=writer)
        outcome = exporter.export([_record("https://e.com/a.html")], work_dir=tmp_path)
        assert outcome.file_path is None
        assert outcome.error.error_code == ErrorCode.EXPORT_WRITE_ERROR

    async def test_url_integrity(self, exporter, tmp_path):
        url = "https://support.huaweicloud.com/intl/zh-cn/usermanual-cli/codeartsagent_cli_0001.html"
        records = [_record(url, "标题")]
        outcome = exporter.export(records, work_dir=tmp_path)
        wb = load_workbook(outcome.file_path)
        ws = wb.active
        assert ws.cell(2, 1).value == url

    async def test_row_order_preserved(self, exporter, tmp_path):
        records = [
            _record("https://e.com/c.html", "C"),
            _record("https://e.com/a.html", "A"),
            _record("https://e.com/b.html", "B"),
        ]
        outcome = exporter.export(records, work_dir=tmp_path)
        wb = load_workbook(outcome.file_path)
        ws = wb.active
        assert ws.cell(2, 1).value == "https://e.com/c.html"
        assert ws.cell(3, 1).value == "https://e.com/a.html"
        assert ws.cell(4, 1).value == "https://e.com/b.html"

    async def test_filename_fixed(self, exporter, tmp_path):
        outcome = exporter.export([], work_dir=tmp_path)
        assert outcome.file_path.endswith(EXPORT_FILE_SPEC.file_name)
