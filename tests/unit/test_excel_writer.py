"""ExcelWriter 单元测试（T8.02）。"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from url_link_extractor.exceptions import ExcelWriteError
from url_link_extractor.infra.excel_writer import ExcelWriter
from url_link_extractor.models.export_spec import EXPORT_FILE_SPEC


@pytest.fixture
def writer():
    return ExcelWriter()


@pytest.fixture
def tmp_workdir(tmp_path):
    return tmp_path


class TestWrite:
    def test_normal_write(self, writer, tmp_workdir):
        path = tmp_workdir / "result.xlsx"
        rows = [
            ("https://example.com/a.html", "标题A"),
            ("https://example.com/b.html", "标题B"),
        ]
        writer.write(rows, path)
        wb = load_workbook(str(path))
        ws = wb.active
        assert ws.cell(1, 1).value == "URL 完整链接"
        assert ws.cell(1, 2).value == "页面标题"
        assert ws.cell(2, 1).value == "https://example.com/a.html"
        assert ws.cell(2, 2).value == "标题A"
        assert ws.cell(3, 1).value == "https://example.com/b.html"
        assert ws.cell(3, 2).value == "标题B"

    def test_empty_rows_only_header(self, writer, tmp_workdir):
        path = tmp_workdir / "result.xlsx"
        writer.write([], path)
        wb = load_workbook(str(path))
        ws = wb.active
        assert ws.cell(1, 1).value == "URL 完整链接"
        assert ws.cell(1, 2).value == "页面标题"
        assert ws.max_row == 1

    def test_overwrite_old_file(self, writer, tmp_workdir):
        path = tmp_workdir / "result.xlsx"
        writer.write([("https://e.com/old.html", "旧")], path)
        writer.write([("https://e.com/new.html", "新")], path)
        wb = load_workbook(str(path))
        ws = wb.active
        assert ws.max_row == 2
        assert ws.cell(2, 1).value == "https://e.com/new.html"
        assert ws.cell(2, 2).value == "新"

    def test_title_none_as_empty(self, writer, tmp_workdir):
        path = tmp_workdir / "result.xlsx"
        writer.write([("https://e.com/a.html", None)], path)
        wb = load_workbook(str(path))
        ws = wb.active
        # openpyxl 将空字符串转为 None，均表示缺失
        assert ws.cell(2, 2).value in (None, "")

    def test_title_empty_string(self, writer, tmp_workdir):
        path = tmp_workdir / "result.xlsx"
        writer.write([("https://e.com/a.html", "")], path)
        wb = load_workbook(str(path))
        ws = wb.active
        assert ws.cell(2, 2).value in (None, "")

    def test_url_integrity_preserved(self, writer, tmp_workdir):
        path = tmp_workdir / "result.xlsx"
        url = "https://support.huaweicloud.com/intl/zh-cn/usermanual-cli/codeartsagent_cli_0001.html"
        writer.write([(url, "标题")], path)
        wb = load_workbook(str(path))
        ws = wb.active
        assert ws.cell(2, 1).value == url

    def test_filename_and_columns_fixed(self, writer, tmp_workdir):
        path = tmp_workdir / EXPORT_FILE_SPEC.file_name
        writer.write([], path)
        assert path.exists()
        wb = load_workbook(str(path))
        ws = wb.active
        assert ws.cell(1, 1).value == EXPORT_FILE_SPEC.column_1_name
        assert ws.cell(1, 2).value == EXPORT_FILE_SPEC.column_2_name
        assert ws.max_column == EXPORT_FILE_SPEC.column_count

    def test_workdir_not_writable(self, writer, tmp_path):
        from unittest.mock import patch

        path = tmp_path / "result.xlsx"
        with patch("os.access", return_value=False):
            with pytest.raises(ExcelWriteError):
                writer.write([], path)
