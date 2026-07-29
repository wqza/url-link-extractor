"""result.xlsx 端到端验证（T15.02）。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from url_link_extractor.core.result_exporter import ExcelResultExporter
from url_link_extractor.models.enums import ErrorCode, RecordStatus
from url_link_extractor.models.export_spec import EXPORT_FILE_SPEC
from url_link_extractor.models.result import LinkTitleRecord

PREFIX = "https://support.huaweicloud.com/intl/zh-cn/usermanual-cli/"


def _record(url, title="T", status=RecordStatus.SUCCESS):
    return LinkTitleRecord(url=url, title=title, status=status)


class TestExportE2E:
    async def test_export_and_verify(self, tmp_path):
        records = [
            _record(PREFIX + "codeartsagent_cli_0001.html", "什么是码道CLI"),
            _record(PREFIX + "codeartsagent_cli_0002.html", title=None, status=RecordStatus.NO_TITLE),
            _record(PREFIX + "codeartsagent_cli_0003.html", title=None, status=RecordStatus.FAILED),
        ]
        exporter = ExcelResultExporter()
        outcome = exporter.export(records, work_dir=tmp_path)
        assert outcome.error is None

        wb = load_workbook(outcome.file_path)
        ws = wb.active

        # 文件名固定
        assert Path(outcome.file_path).name == EXPORT_FILE_SPEC.file_name

        # 列名固定、列数 2
        assert ws.cell(1, 1).value == EXPORT_FILE_SPEC.column_1_name
        assert ws.cell(1, 2).value == EXPORT_FILE_SPEC.column_2_name
        assert ws.max_column == EXPORT_FILE_SPEC.column_count

        # 行数 = 结果集大小
        assert ws.max_row - 1 == len(records)

        # URL 列绝对完整，以前缀开头
        for i, record in enumerate(records, start=2):
            url_cell = ws.cell(i, 1).value
            assert url_cell.startswith("https://")
            assert url_cell.startswith(PREFIX)
            assert "..." not in url_cell
            assert "./" not in url_cell

        # 标题缺失行第二列为空
        assert ws.cell(3, 2).value in (None, "")
        assert ws.cell(4, 2).value in (None, "")

        # 行顺序与结果集一致
        assert ws.cell(2, 1).value == records[0].url

    async def test_empty_result_export(self, tmp_path):
        exporter = ExcelResultExporter()
        outcome = exporter.export([], work_dir=tmp_path)
        assert outcome.error is None
        wb = load_workbook(outcome.file_path)
        ws = wb.active
        assert ws.max_row == 1  # 仅表头
        assert ws.cell(1, 1).value == EXPORT_FILE_SPEC.column_1_name

    async def test_overwrite(self, tmp_path):
        exporter = ExcelResultExporter()
        exporter.export([_record(PREFIX + "old.html", "旧")], work_dir=tmp_path)
        exporter.export([_record(PREFIX + "new.html", "新")], work_dir=tmp_path)
        wb = load_workbook(str(tmp_path / "result.xlsx"))
        ws = wb.active
        assert ws.max_row == 2
        assert ws.cell(2, 1).value == PREFIX + "new.html"

    async def test_export_failure_soft(self, tmp_path):
        from unittest.mock import MagicMock

        from url_link_extractor.exceptions import ExcelWriteError

        writer = MagicMock()
        writer.write = MagicMock(side_effect=ExcelWriteError("disk full"))
        exporter = ExcelResultExporter(writer=writer)
        outcome = exporter.export([_record(PREFIX + "a.html")], work_dir=tmp_path)
        assert outcome.file_path is None
        assert outcome.error.error_code == ErrorCode.EXPORT_WRITE_ERROR
